from __future__ import annotations

import os
import shutil
import time

import pytest
from conftest import MENINGIOMA_TEXT, TOKEN, VS_TEXT, make_scanned_pdf, make_textbook

from app.services import extractor, jobs


def scan(client, **body):
    r = client.post("/api/scan", json=body, headers=TOKEN)
    assert r.status_code == 202, r.text
    job = jobs.wait(r.json()["job_id"], timeout=300)
    assert job["status"] == "done", job["log"]
    return job


def add_folder(client, path):
    r = client.post("/api/folders", json={"path": str(path)}, headers=TOKEN)
    assert r.status_code == 201, r.text
    return r.json()


def docs(client, **params):
    return client.get("/api/documents", params=params, headers=TOKEN).json()["items"]


def search(client, q, **params):
    r = client.get("/api/search", params={"q": q, **params}, headers=TOKEN)
    assert r.status_code == 200, r.text
    return r.json()


@pytest.fixture()
def indexed(library, client):
    # VS text on physical page index 5 (printed p. 2), meningioma text on index 15 (printed p. 12)
    make_textbook(library / "Greenberg Handbook of Neurosurgery 9th ed.pdf",
                  extra_page_text={5: VS_TEXT, 15: MENINGIOMA_TEXT})
    folder = add_folder(client, library)
    scan(client)
    return library, folder


def test_auth_required(client, library):
    assert client.get("/api/stats").status_code == 401
    assert client.get("/api/stats", headers={"X-Academy-Token": "wrong"}).status_code == 401
    assert client.get("/api/health").status_code == 200


def test_add_folder_rejects_missing_path(client, library):
    r = client.post("/api/folders", json={"path": str(library / "nope")}, headers=TOKEN)
    assert r.status_code == 400


def test_index_metadata_and_embedded_labels(client, indexed):
    (d,) = docs(client)
    assert d["status"] == "indexed"
    assert d["page_count"] == 30 and d["label_mode"] == "embedded" and d["has_toc"] == 1
    assert d["edition"] == "9th" and d["year"] == 2019 and d["publisher"] == "Thieme"
    assert d["title"] == "Handbook of Test Neurosurgery" and d["title_src"] == "pdf_metadata"
    assert d["metadata_verified"] == 0


def test_search_returns_exact_page_and_chapter(client, indexed):
    res = search(client, "retrosigmoid hearing preservation")
    hit = res["hits"][0]
    c = hit["citation"]
    assert hit["page_index"] == 5
    assert c["pdf_page"] == 6 and c["printed_page"] == "2" and c["printed_page_source"] == "embedded"
    assert c["chapter"] == "Chapter 1. Skull Base Tumors"
    assert c["section"] == "Vestibular Schwannoma"
    assert "p. 2 [PDF p. 6]" in c["formatted"] and "9th ed." in c["formatted"]


def test_synonym_search_finds_textbook_term(client, indexed):
    res = search(client, "acoustic neuroma brainstem compression")
    assert res["hits"] and res["hits"][0]["page_index"] == 5


def test_separate_topics_hit_separate_pages(client, indexed):
    res = search(client, "petroclival meningioma basilar perforators")
    assert res["hits"][0]["page_index"] == 15
    assert res["hits"][0]["citation"]["chapter"] == "Chapter 2. Meningiomas"


def test_no_hit_reports_ungrounded(client, indexed):
    res = search(client, "zzqx nonexistentterm")
    assert res["hits"] == [] and "cannot be grounded" in res["coverage_note"]


def test_citation_verification(client, indexed):
    (d,) = docs(client)
    ok = client.post("/api/citations/verify", headers=TOKEN, json={
        "document_id": d["id"], "page_index": 5,
        "quote": "the translabyrinthine approach sacrifices hearing"}).json()
    assert ok["verified"] and ok["match"] == "exact"
    fuzzy = client.post("/api/citations/verify", headers=TOKEN, json={
        "document_id": d["id"], "page_index": 5,
        "quote": "the translabyrinthine aproach sacrifices hearing"}).json()
    assert fuzzy["verified"] and fuzzy["match"] == "fuzzy"
    off_by_one = client.post("/api/citations/verify", headers=TOKEN, json={
        "document_id": d["id"], "page_index": 6, "quote": "the translabyrinthine approach sacrifices hearing"}).json()
    assert not off_by_one["verified"] and off_by_one["page_index"] == 5
    fabricated = client.post("/api/citations/verify", headers=TOKEN, json={
        "document_id": d["id"], "page_index": 5, "quote": "radiosurgery is always superior to microsurgery"}).json()
    assert not fabricated["verified"]
    unindexed = client.post("/api/citations/verify", headers=TOKEN, json={
        "document_id": 9999, "page_index": 0, "quote": "anything at all here"}).json()
    assert not unindexed["verified"] and "not indexed" in unindexed["reason"]


def test_incremental_rescan_skips_unchanged(client, indexed):
    job = scan(client)
    assert job["progress"]["files_unchanged"] == 1
    assert job["progress"]["to_index"] == 0


def test_exact_duplicate_detected_not_reindexed(client, indexed):
    lib, _ = indexed
    sub = lib / "Telegram Desktop"
    sub.mkdir()
    shutil.copy(lib / "Greenberg Handbook of Neurosurgery 9th ed.pdf", sub / "copy (1).pdf")
    job = scan(client)
    assert job["progress"]["files_duplicates"] == 1
    dup = [d for d in docs(client) if d["status"] == "duplicate"]
    assert len(dup) == 1 and dup[0]["duplicate_of"] is not None
    # duplicates never produce double hits
    res = search(client, "retrosigmoid hearing preservation")
    assert len({h["document_id"] for h in res["hits"]}) == 1


def test_changed_file_is_reindexed(client, indexed):
    lib, _ = indexed
    p = lib / "Greenberg Handbook of Neurosurgery 9th ed.pdf"
    time.sleep(1.1)
    make_textbook(p, extra_page_text={7: "Telovelar approach to the fourth ventricle avoids vermian splitting."})
    os.utime(p, (time.time() + 5, time.time() + 5))
    job = scan(client)
    assert job["progress"]["files_changed"] == 1
    assert search(client, "telovelar vermian")["hits"][0]["page_index"] == 7
    assert search(client, "retrosigmoid hearing preservation")["hits"] == []


def test_deleted_file_marked_missing_and_removed_from_search(client, indexed):
    lib, _ = indexed
    (lib / "Greenberg Handbook of Neurosurgery 9th ed.pdf").unlink()
    job = scan(client)
    assert job["progress"]["files_missing"] == 1
    assert docs(client)[0]["status"] == "missing"
    assert search(client, "retrosigmoid hearing preservation")["hits"] == []


def test_moved_file_keeps_index_without_reindex(client, indexed):
    lib, _ = indexed
    (lib / "moved").mkdir()
    shutil.move(lib / "Greenberg Handbook of Neurosurgery 9th ed.pdf", lib / "moved" / "Greenberg.pdf")
    job = scan(client)
    assert job["progress"]["files_moved"] == 1 and job["progress"]["to_index"] == 0
    (d,) = docs(client)
    assert d["path"].endswith("Greenberg.pdf") and d["status"] == "indexed"
    assert search(client, "retrosigmoid hearing preservation")["hits"]


def test_duplicate_promoted_when_original_deleted(client, indexed):
    lib, _ = indexed
    shutil.copy(lib / "Greenberg Handbook of Neurosurgery 9th ed.pdf", lib / "backup.pdf")
    scan(client)
    (lib / "Greenberg Handbook of Neurosurgery 9th ed.pdf").unlink()
    scan(client)
    live = [d for d in docs(client) if d["status"] == "indexed"]
    assert len(live) == 1 and live[0]["filename"] == "backup.pdf"
    assert search(client, "retrosigmoid hearing preservation")["hits"][0]["document_id"] == live[0]["id"]


def test_unreachable_folder_does_not_mark_missing(client, indexed):
    lib, _ = indexed
    hidden = lib.parent / "offline"
    lib.rename(hidden)  # simulates an unplugged drive / offline share
    job = scan(client)
    assert job["progress"].get("files_missing", 0) == 0
    assert docs(client)[0]["status"] == "indexed"
    hidden.rename(lib)


def test_inferred_page_numbers_flagged(client, library):
    make_textbook(library / "no labels.pdf", embedded_labels=False, extra_page_text={9: VS_TEXT})
    add_folder(client, library)
    scan(client)
    (d,) = docs(client)
    assert d["label_mode"] == "inferred"
    c = search(client, "retrosigmoid hearing preservation")["hits"][0]["citation"]
    assert c["printed_page"] == "6" and c["printed_page_source"] == "inferred"
    assert "inferred" in c["formatted"]


def test_no_page_numbers_cites_pdf_page_only(client, library):
    make_textbook(library / "bare.pdf", embedded_labels=False, footer_numbers=False, extra_page_text={9: VS_TEXT})
    add_folder(client, library)
    scan(client)
    c = search(client, "retrosigmoid hearing preservation")["hits"][0]["citation"]
    assert c["printed_page"] is None and "PDF p. 10 (printed page number unavailable)" in c["formatted"]


def test_reflowable_formats(client, library):
    (library / "notes.md").write_text("# Pineal region\n\nThe supracerebellar infratentorial approach uses gravity.\n",
                                      encoding="utf-8")
    import docx
    d = docx.Document()
    d.add_heading("Cavernous sinus", level=1)
    d.add_paragraph("The abducens nerve lies free within the cavernous sinus lateral to the ICA.")
    d.save(str(library / "lecture.docx"))
    add_folder(client, library)
    scan(client)
    assert {x["status"] for x in docs(client)} == {"indexed"}
    c = search(client, "abducens cavernous")["hits"][0]["citation"]
    assert c["reflowable"] and "Section 1" in c["formatted"] and c["chapter"] == "Cavernous sinus"
    assert search(client, "supracerebellar infratentorial gravity")["hits"]


@pytest.mark.skipif(not shutil.which("tesseract"), reason="tesseract not installed")
def test_ocr_scanned_pdf(client, library):
    make_scanned_pdf(library / "scanned.pdf", "Chiari malformation decompression with duraplasty " * 4)
    add_folder(client, library)
    scan(client)
    (d,) = docs(client)
    assert d["ocr_pages"] == 1
    hit = search(client, "chiari duraplasty")["hits"][0]
    assert hit["citation"]["ocr"] is True


def test_scanned_pdf_without_ocr_flagged(client, library, monkeypatch):
    monkeypatch.setattr(extractor, "ocr_available", lambda: False)
    make_scanned_pdf(library / "scanned.pdf", "Chiari malformation")
    add_folder(client, library)
    scan(client)
    (d,) = docs(client)
    assert d["status"] == "ocr_needed" and "not OCRed" in d["quality_notes"]


def test_user_metadata_survives_reindex(client, indexed):
    (d,) = docs(client)
    r = client.patch(f"/api/documents/{d['id']}", headers=TOKEN,
                     json={"authors": "Greenberg MS", "metadata_verified": True})
    assert r.json()["authors_src"] == "user"
    scan(client, force_reindex=True)
    (d2,) = docs(client)
    assert d2["authors"] == "Greenberg MS" and d2["metadata_verified"] == 1
    assert search(client, "retrosigmoid hearing", verified_only=True)["hits"]


def test_page_image_file_and_printed_page_lookup(client, indexed):
    (d,) = docs(client)
    img = client.get(f"/api/documents/{d['id']}/pages/5/image", params={"highlight": "retrosigmoid"}, headers=TOKEN)
    assert img.status_code == 200 and img.content[:8] == b"\x89PNG\r\n\x1a\n"
    f = client.get(f"/api/documents/{d['id']}/file", headers=TOKEN)
    assert f.status_code == 200 and f.content[:4] == b"%PDF"
    m = client.get(f"/api/documents/{d['id']}/resolve-page", params={"printed": "2"}, headers=TOKEN).json()
    assert m["matches"][0]["page_index"] == 5
    page = client.get(f"/api/documents/{d['id']}/pages/5", headers=TOKEN).json()
    assert "retrosigmoid" in page["text"]


def test_original_file_unchanged_after_index_and_render(client, indexed):
    lib, _ = indexed
    p = lib / "Greenberg Handbook of Neurosurgery 9th ed.pdf"
    from app.services.indexer import sha256_file
    before = sha256_file(str(p))
    (d,) = docs(client)
    client.get(f"/api/documents/{d['id']}/pages/5/image", params={"highlight": "retrosigmoid"}, headers=TOKEN)
    scan(client, force_reindex=True)
    assert sha256_file(str(p)) == before


def test_concurrent_scan_rejected(client, indexed):
    lib, _ = indexed
    for i in range(3):
        make_textbook(lib / f"extra{i}.pdf", pages=120, title=f"Extra {i}")
    r1 = client.post("/api/scan", json={}, headers=TOKEN)
    r2 = client.post("/api/scan", json={}, headers=TOKEN)
    assert r2.status_code == 409
    jobs.wait(r1.json()["job_id"])


def test_stats_and_quality_report(client, indexed):
    s = client.get("/api/stats", headers=TOKEN).json()
    assert s["indexed_documents"] == 1 and s["pages"] == 30 and s["questions_generated"] is None
    q = client.get("/api/quality-report", headers=TOKEN).json()
    assert q[0]["quality_score"] >= 90


def test_remove_folder_deindexes_without_touching_files(client, indexed):
    lib, folder = indexed
    r = client.delete(f"/api/folders/{folder['id']}", headers=TOKEN)
    assert r.json()["removed_documents"] == 1
    assert (lib / "Greenberg Handbook of Neurosurgery 9th ed.pdf").exists()
    assert search(client, "retrosigmoid hearing")["hits"] == []


def test_scan_blocked_while_other_process_holds_lock(client, indexed):
    from app.services.lock import ScanLock
    other = ScanLock()
    assert other.acquire()
    try:
        r = client.post("/api/scan", json={}, headers=TOKEN)
        assert r.status_code == 409 and "another indexing process" in r.json()["detail"]
    finally:
        other.release()
    scan(client)
