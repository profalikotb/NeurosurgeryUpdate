"""Text extraction. Originals are opened read-only and never modified."""
from __future__ import annotations

import logging
import os
import re
import shutil
from collections.abc import Callable
from dataclasses import dataclass, field

import pymupdf

from ..config import get_settings
from .paths import os_path

log = logging.getLogger(__name__)

EXTRACTOR_VERSION = 2  # bump to trigger re-indexing of existing libraries after extraction changes
REFLOWABLE_EXT = {".epub", ".mobi", ".fb2"}
PSEUDO_PAGE_CHARS = 3000  # DOCX/TXT have no physical pages; they are split into numbered sections


@dataclass
class ExtractedPage:
    text: str
    ocr: bool = False
    has_images: bool = False


@dataclass
class ExtractedDoc:
    pages: list[ExtractedPage]
    toc: list[tuple[int, str, int]] = field(default_factory=list)  # (level, title, 0-based start page)
    embedded_labels: list[str] = field(default_factory=list)
    embedded_meta: dict[str, str] = field(default_factory=dict)
    reflowable: bool = False
    ocr_skipped_pages: int = 0
    warnings: list[str] = field(default_factory=list)


_tesseract_checked: bool | None = None


def ocr_available() -> bool:
    global _tesseract_checked
    if _tesseract_checked is None:
        s = get_settings()
        if s.tessdata_prefix:
            os.environ.setdefault("TESSDATA_PREFIX", s.tessdata_prefix)
        _tesseract_checked = bool(shutil.which("tesseract")) or bool(os.environ.get("TESSDATA_PREFIX"))
    return _tesseract_checked and get_settings().ocr_enabled


def clean_text(t: str) -> str:
    t = t.replace("­", "")                       # soft hyphen
    t = re.sub(r"(\w)-\n(\w)", r"\1\2", t)            # residual hyphenation at line end
    t = re.sub(r"[ \t ]+", " ", t)
    t = re.sub(r"\n{3,}", "\n\n", t)
    return t.strip()


def extract(path: str, ext: str, progress: Callable[[int, int], None] | None = None,
            cancelled: Callable[[], bool] | None = None) -> ExtractedDoc:
    ext = ext.lower()
    if ext == ".docx":
        return _extract_docx(path)
    if ext in (".txt", ".md"):
        return _extract_text(path)
    return _extract_mupdf(path, ext, progress, cancelled)


def _extract_mupdf(path: str, ext: str, progress, cancelled) -> ExtractedDoc:
    s = get_settings()
    doc = pymupdf.open(os_path(path))
    try:
        if doc.needs_pass:
            raise ValueError("document is password protected")
        out = ExtractedDoc(pages=[], reflowable=ext in REFLOWABLE_EXT or doc.is_reflowable)
        out.embedded_meta = {k: (v or "") for k, v in (doc.metadata or {}).items()}
        try:
            out.toc = [(lvl, title.strip(), max(0, pg - 1)) for lvl, title, pg in doc.get_toc(simple=True)
                       if title and title.strip() and pg >= 1]
        except Exception as e:  # malformed outlines are common in scanned books
            out.warnings.append(f"TOC unreadable: {e}")

        can_ocr = ocr_available()
        ocr_budget = s.ocr_max_pages_per_document
        flags = pymupdf.TEXT_DEHYPHENATE | pymupdf.TEXT_PRESERVE_WHITESPACE | pymupdf.TEXT_MEDIABOX_CLIP
        n = doc.page_count
        for i in range(n):
            if cancelled and cancelled():
                raise InterruptedError("cancelled")
            page = doc.load_page(i)
            label = ""
            if not out.reflowable:
                try:
                    label = page.get_label() or ""
                except Exception:
                    label = ""
            out.embedded_labels.append(label)
            text = page.get_text("text", flags=flags) or ""
            has_images = bool(page.get_images(full=False))
            ep = ExtractedPage(text=clean_text(text), has_images=has_images)
            if len(ep.text) < s.min_text_chars_per_page and has_images:
                if can_ocr and ocr_budget > 0:
                    try:
                        tp = page.get_textpage_ocr(language=s.ocr_language, dpi=s.ocr_dpi, full=True)
                        ocr_text = clean_text(page.get_text("text", textpage=tp) or "")
                        if len(ocr_text) > len(ep.text):
                            ep = ExtractedPage(text=ocr_text, ocr=True, has_images=True)
                        ocr_budget -= 1
                    except Exception as e:
                        out.warnings.append(f"OCR failed on PDF p.{i + 1}: {e}")
                        can_ocr = False
                else:
                    out.ocr_skipped_pages += 1
            out.pages.append(ep)
            if progress and (i % 10 == 0 or i == n - 1):
                progress(i + 1, n)
        return out
    finally:
        doc.close()


def _pseudo_pages(blocks: list[tuple[str, int | None]]) -> ExtractedDoc:
    """blocks: (text, heading_level or None). Builds ~3000-char sections and a TOC from headings."""
    pages: list[ExtractedPage] = []
    toc: list[tuple[int, str, int]] = []
    buf: list[str] = []
    size = 0
    for text, level in blocks:
        if level is not None and buf and size > PSEUDO_PAGE_CHARS * 0.3:
            pages.append(ExtractedPage(clean_text("\n".join(buf))))
            buf, size = [], 0
        if level is not None:
            toc.append((level, text.strip(), len(pages)))
        buf.append(text)
        size += len(text)
        if size >= PSEUDO_PAGE_CHARS:
            pages.append(ExtractedPage(clean_text("\n".join(buf))))
            buf, size = [], 0
    if buf:
        pages.append(ExtractedPage(clean_text("\n".join(buf))))
    return ExtractedDoc(pages=pages, toc=toc, reflowable=True)


def _extract_docx(path: str) -> ExtractedDoc:
    import docx  # python-docx

    d = docx.Document(os_path(path))
    blocks: list[tuple[str, int | None]] = []
    for p in d.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        level = None
        m = re.match(r"heading\s*(\d)", style)
        if m:
            level = int(m.group(1))
        elif style == "title":
            level = 1
        blocks.append((txt, level))
    for t in d.tables:
        rows = [" | ".join(c.text.strip() for c in r.cells) for r in t.rows]
        blocks.append(("\n".join(rows), None))
    out = _pseudo_pages(blocks)
    cp = d.core_properties
    out.embedded_meta = {"title": cp.title or "", "author": cp.author or ""}
    return out


def _extract_text(path: str) -> ExtractedDoc:
    with open(os_path(path), "rb") as f:
        raw = f.read()
    for enc in ("utf-8-sig", "utf-16", "cp1256", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    blocks: list[tuple[str, int | None]] = []
    for para in re.split(r"\n\s*\n", text):
        m = re.match(r"^(#{1,4})\s+(.*)", para.strip())
        if m:
            blocks.append((m.group(2), len(m.group(1))))
        elif para.strip():
            blocks.append((para, None))
    return _pseudo_pages(blocks)
